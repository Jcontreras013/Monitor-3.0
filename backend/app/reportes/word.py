"""
Formulario oficial de Recursos Humanos "REPORTE DE FALTAS" (HN-GG-RH-FR-09)
en Word, como lo generaba el sistema anterior. Con varios colaboradores sale
un formulario por persona, cada uno en su página.
"""

from __future__ import annotations

import io
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

from app.reportes import ruta_logo

AZUL = "1A1D54"
FICHA = (("CÓDIGO", "HN-GG-RH-FR-09"), ("VERSIÓN", "1.0"), ("FECHA", "12/05/2026"), ("CLASIFICACIÓN", "INTERNO"))
FILAS_DETALLE = 12


@dataclass
class Incidencia:
    fecha: date
    motivo: str
    comentario: str
    hora: str


@dataclass
class FichaEmpleado:
    nombre: str
    puesto: str
    departamento: str
    incidencias: Sequence[Incidencia]


def _sombrear(celda, color: str) -> None:
    celda._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>'))


def _texto(celda, texto: str, *, negrita=False, blanco=False, tamano=8.5, centro=False) -> None:
    p = celda.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centro else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(1.5)
    p.paragraph_format.space_after = Pt(1.5)
    run = p.add_run(texto)
    run.bold = negrita
    run.font.size = Pt(tamano)
    if blanco:
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def _seccion(doc, titulo: str, filas: Sequence[tuple[str, str]]) -> None:
    tabla = doc.add_table(rows=len(filas) + 1, cols=2)
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    encabezado = tabla.rows[0].cells[0].merge(tabla.rows[0].cells[1])
    _sombrear(encabezado, AZUL)
    _texto(encabezado, titulo, negrita=True, blanco=True)
    for i, (etiqueta, valor) in enumerate(filas, start=1):
        celda_etq, celda_val = tabla.rows[i].cells
        celda_etq.width, celda_val.width = Inches(2.6), Inches(4.9)
        _sombrear(celda_etq, "E8EAF2")
        _texto(celda_etq, etiqueta, negrita=True)
        _texto(celda_val, valor)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _formulario(doc, ficha: FichaEmpleado, ahora: datetime) -> None:
    # Encabezado: logo | título | ficha del documento
    enc = doc.add_table(rows=1, cols=3)
    enc.style = "Table Grid"
    c_logo, c_titulo, c_ficha = enc.rows[0].cells
    c_logo.width, c_titulo.width, c_ficha.width = Inches(1.8), Inches(3.7), Inches(2.0)
    c_logo.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c_logo.paragraphs[0].add_run().add_picture(str(ruta_logo()), width=Inches(1.5))
    _texto(c_titulo, "REPORTE DE FALTAS", negrita=True, tamano=14, centro=True)
    for i, (k, v) in enumerate(FICHA):
        p = c_ficha.paragraphs[0] if i == 0 else c_ficha.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(f"{k}: ")
        r.bold, r.font.size = True, Pt(7.5)
        p.add_run(v).font.size = Pt(7.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    fechas = sorted({i.fecha for i in ficha.incidencias})
    horas = [i.hora for i in ficha.incidencias if i.hora and i.hora != "N/D"]
    _seccion(doc, "Datos del Jefe del Departamento Solicitante", [
        ("Nombre del Jefe del Departamento", ""), ("Puesto", ""), ("Departamento", ""),
        ("Ciudad", ""), ("Fecha de Ingreso", ""),
    ])
    _seccion(doc, "Datos de Empleado Reportado", [
        ("Nombre del Empleado", ficha.nombre),
        ("Código de Empleado", ""),
        ("Puesto", ficha.puesto),
        ("Horario de Trabajo", ""),
        ("Fecha de la Falta Ocurrida", " - ".join(f.strftime("%d/%m/%Y") for f in fechas) or "N/D"),
        ("Hora de la Falta Ocurrida", " - ".join(horas) or "N/D"),
    ])

    detalle = doc.add_table(rows=1 + max(FILAS_DETALLE, len(ficha.incidencias)), cols=1)
    detalle.style = "Table Grid"
    _sombrear(detalle.rows[0].cells[0], AZUL)
    _texto(detalle.rows[0].cells[0], "Detalle de la Falta", negrita=True, blanco=True)
    for i, inc in enumerate(sorted(ficha.incidencias, key=lambda x: x.fecha), start=1):
        _texto(detalle.rows[i].cells[0], f"• [{inc.fecha:%d/%m/%Y}] {inc.motivo}: {inc.comentario}")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.add_run("Firmo en señal de solicitud/autorización/aprobación lo de arriba detallado.").italic = True

    firmas = doc.add_table(rows=2, cols=2)
    firmas.style = "Table Grid"
    for celda, titulo in zip(firmas.rows[0].cells, ("Firma de Jefe de\nDepartamento Solicitante",
                                                    "Firma de Recursos Humanos"), strict=True):
        _sombrear(celda, AZUL)
        _texto(celda, titulo, negrita=True, blanco=True, centro=True)
    for celda in firmas.rows[1].cells:
        celda.paragraphs[0].add_run("\n\n\n")
    pie = doc.add_paragraph()
    pie.paragraph_format.space_before = Pt(4)
    texto_pie = f"Elaborado por el sistema Monitor Operativo el {ahora:%d/%m/%Y} a las {ahora:%I:%M %p}"
    pie.add_run(texto_pie).font.size = Pt(7)


def reporte_faltas(fichas: Sequence[FichaEmpleado], ahora: datetime) -> bytes:
    doc = Document()
    for seccion in doc.sections:
        seccion.top_margin = seccion.bottom_margin = Inches(0.5)
        seccion.left_margin = seccion.right_margin = Inches(0.5)
    estilo = doc.styles["Normal"]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(8.5)

    for i, ficha in enumerate(fichas):
        if i:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _formulario(doc, ficha, ahora)
    if not fichas:
        doc.add_paragraph("No hay registros para los filtros elegidos.")

    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()
