"""
Clientes VIP: lectura del Excel/CSV que entrega el área comercial.

Del archivo solo se toman código, nombre y clase (el original trae correos y
teléfonos de contacto, que no se guardan). Port de clientes_vip.py.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass

COLUMNAS_NOMBRE = ("NOMBRE", "NOMBRE CLIENTE", "RAZON SOCIAL")
COLUMNAS_CLASE = ("CLASE CLIENTE", "CLASE")
_CODIGO = re.compile(r"\d+")


@dataclass(frozen=True)
class ClienteVip:
    cliente: str
    nombre: str = ""
    clase: str = ""


def normalizar_codigo(valor) -> str:
    """Código de cliente como texto sin '.0' (Excel lo devuelve como número)."""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = "" if valor is None else str(valor).strip()
    return texto[:-2] if texto.endswith(".0") else texto


def _texto(valor) -> str:
    texto = "" if valor is None else str(valor).strip()
    return "" if texto.lower() in ("nan", "none") else texto


def _desde_filas(filas: list[list]) -> list[ClienteVip] | None:
    """Usa la primera fila como encabezado; None si no hay columna CLIENTE."""
    if not filas:
        return None
    columnas = {str(c).strip().upper(): i for i, c in enumerate(filas[0]) if c is not None}
    if "CLIENTE" not in columnas:
        return None
    i_cliente = columnas["CLIENTE"]
    i_nombre = next((columnas[c] for c in COLUMNAS_NOMBRE if c in columnas), None)
    i_clase = next((columnas[c] for c in COLUMNAS_CLASE if c in columnas), None)
    vistos: dict[str, ClienteVip] = {}
    for fila in filas[1:]:
        celda = lambda i: fila[i] if i is not None and i < len(fila) else None  # noqa: E731
        codigo = normalizar_codigo(celda(i_cliente))
        if _CODIGO.fullmatch(codigo) and codigo not in vistos:
            vistos[codigo] = ClienteVip(codigo, _texto(celda(i_nombre)), _texto(celda(i_clase)))
    return list(vistos.values())


def leer_archivo_vip(contenido: bytes, nombre_archivo: str) -> list[ClienteVip]:
    """
    Lee la lista de un .xlsx o .csv. Busca la hoja que tenga una columna
    CLIENTE (el archivo puede traer hojas vacías o con otro nombre).
    """
    nombre = nombre_archivo.lower()
    if nombre.endswith(".csv"):
        try:
            texto = contenido.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = contenido.decode("latin-1")
        hojas = [list(csv.reader(io.StringIO(texto)))]
    elif nombre.endswith(".xlsx"):
        from openpyxl import load_workbook

        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        hojas = [[list(f) for f in hoja.iter_rows(values_only=True)] for hoja in libro.worksheets]
        libro.close()
    else:
        raise ValueError("Sube el archivo como Excel (.xlsx) o CSV. Si es un .xls antiguo, guárdalo como .xlsx.")
    for filas in hojas:
        clientes = _desde_filas(filas)
        if clientes is not None:
            return clientes
    raise ValueError("El archivo no tiene una columna llamada CLIENTE con los códigos de cliente.")
