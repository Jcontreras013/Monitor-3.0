import io
from datetime import timedelta

from openpyxl import load_workbook

from app import repositorio
from app.dominio.fechas import ahora_hn
from app.dominio.normalizacion import normalizar_orden
from tests.utiles import entrar

FMT = "%Y-%m-%d %H:%M:%S"


async def _cargar(sesiones):
    ahora = ahora_hn()
    hoy_8 = ahora.replace(hour=8, minute=0, second=0, microsecond=0)
    hora = lambda dt: dt.strftime(FMT)  # noqa: E731
    crudas = [
        # Cerrada en la jornada de hoy (HORA_LIQ adelantada 6 h, como llega de Cepheus).
        {"NUM": "1", "ACTIVIDAD": "INSFIBRA", "ESTADO": "CERRADA", "TECNICO": "Ana", "SECTOR": "TELA",
         "FECHA_APE": hora(hoy_8 - timedelta(days=2)), "HORA_INI": hora(hoy_8),
         "HORA_LIQ": hora(hoy_8 + timedelta(hours=7))},
        {"NUM": "2", "ACTIVIDAD": "SOPFIBRA", "ESTADO": "CERRADA", "TECNICO": "Beto", "SECTOR": "TELA",
         "FECHA_APE": hora(hoy_8), "HORA_INI": hora(hoy_8 + timedelta(minutes=30)),
         "HORA_LIQ": hora(hoy_8 + timedelta(hours=8))},
        # Pendientes: una asignada vieja y una sin asignar.
        {"NUM": "3", "ACTIVIDAD": "SOPFIBRA", "ESTADO": "ASIGNADA", "TECNICO": "Ana",
         "FECHA_APE": hora(ahora - timedelta(days=8)), "COMENTARIO": "ONU OFFLINE"},
        {"NUM": "4", "ACTIVIDAD": "INSFIBRA", "ESTADO": "PENDIENTE", "TECNICO": "",
         "FECHA_APE": hora(ahora - timedelta(hours=3)), "COLONIA": "Centro"},
    ]
    async with sesiones() as s:
        await repositorio.guardar_ordenes(s, [o for c in crudas if (o := normalizar_orden(c, ahora))])
        await s.commit()
    return ahora.date()


async def test_reportes_solo_admin_y_jefe(cliente, crear_usuario):
    await entrar(cliente, crear_usuario, usuario="mon", rol="monitoreo")
    assert (await cliente.get("/api/reportes/cierre")).status_code == 403


async def test_cierre_diario_y_sus_pdf(cliente, crear_usuario, sesiones):
    hoy = await _cargar(sesiones)
    await entrar(cliente, crear_usuario)

    d = (await cliente.get("/api/reportes/cierre", params={"fecha": hoy.isoformat()})).json()
    assert d["total_cerradas"] == 2
    assert [o["num"] for o in d["cerradas"]] == ["1", "2"]
    assert {t["tecnico"]: t["logro_final"] for t in d["eficiencia"]} == {"Ana": 25.0, "Beto": 12.5}
    assert [p["tecnico"] for p in d["primera_orden"]] == ["Ana", "Beto"]
    glob = d["avance"][2]
    assert (glob["segmento"], glob["cerradas"], glob["total"]) == ("GLOBAL", 2, 3)

    for ruta in ("cierre.pdf", "cerradas.pdf", "primera-orden.pdf"):
        r = await cliente.get(f"/api/reportes/{ruta}", params={"fecha": hoy.isoformat()})
        assert r.status_code == 200 and r.content.startswith(b"%PDF"), ruta

    # Un día sin datos también genera su PDF.
    vacio = await cliente.get("/api/reportes/cierre.pdf", params={"fecha": "2020-01-01"})
    assert vacio.status_code == 200 and vacio.content.startswith(b"%PDF")


async def test_promedio_de_arranque(cliente, crear_usuario, sesiones):
    hoy = await _cargar(sesiones)
    await entrar(cliente, crear_usuario)
    d = (await cliente.get("/api/reportes/arranque")).json()
    assert d["hasta"] == hoy.isoformat()
    assert [(f["tecnico"], f["hora_promedio"]) for f in d["filas"]] == [("Ana", "08:00:00"), ("Beto", "08:30:00")]
    assert (await cliente.get("/api/reportes/arranque.pdf")).content.startswith(b"%PDF")
    malo = await cliente.get("/api/reportes/arranque", params={"desde": "2026-05-10", "hasta": "2026-05-01"})
    assert malo.status_code == 400


async def test_pendientes_para_dispatch(cliente, crear_usuario, sesiones):
    await _cargar(sesiones)
    await entrar(cliente, crear_usuario)
    d = (await cliente.get("/api/reportes/pendientes")).json()
    assert d["totales"] == {"asignadas": 1, "sin_asignar": 1, "total": 2}
    assert [o["num"] for o in d["sin_asignar"]] == ["4"]
    assert d["cubetas"][0] == {"nombre": ">= 7 días", "cantidad": 1}

    assert (await cliente.get("/api/reportes/pendientes.pdf")).content.startswith(b"%PDF")
    r = await cliente.get("/api/reportes/pendientes.xlsx")
    hoja = load_workbook(io.BytesIO(r.content))["Pendientes"]
    filas = list(hoja.iter_rows(values_only=True))
    assert filas[0][:3] == ("Días", "Orden", "Cliente")
    assert [f[1] for f in filas[1:]] == ["3", "4"]  # más días primero
    assert filas[2][10] == "SIN ASIGNAR"
